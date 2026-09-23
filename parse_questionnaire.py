"""Parse an incoming security questionnaire (.xlsx or .docx) into a list of
structured questions. Handles the common shapes these come in:

- Excel-based (SIG Lite/Core, CAIQ, vendor-specific): one sheet per section,
  a "Question" column, and usually an empty "Answer"/"Response" column.
- Word-based: numbered or bulleted paragraphs ending in a question mark,
  sometimes inside tables.

This is heuristic by design; questionnaires are not standardized. Review
`--dry-run` output before running the full pipeline on a new template.
"""
from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, asdict
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument

QUESTION_HEADER_HINTS = ["question", "control", "requirement", "item"]
ANSWER_HEADER_HINTS = ["answer", "response", "vendor response", "comments"]


@dataclass
class Question:
    id: str
    sheet: str
    row: int
    text: str
    existing_answer: str = ""
    # Section heading the question sits under (e.g. "Certifications"), used as
    # extra retrieval context for terse rows like "ISO 27001/27002".
    section: str = ""
    # .docx only: column index of the answer cell in a Question/Answer table.
    # None means there's no known answer slot (paragraph or unstructured table).
    answer_col: int | None = None

    @property
    def prompt_text(self) -> str:
        return f"{self.section}: {self.text}" if self.section else self.text


NUMBER_RE = re.compile(r"^\d+(\.\d+)*\.?$")


def _find_column(columns, hints) -> str | None:
    lower = {str(c).strip().lower(): c for c in columns}
    for hint in hints:
        for lc, orig in lower.items():
            if hint in lc:
                return orig
    return None


def parse_xlsx(path: Path) -> list[Question]:
    questions: list[Question] = []
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    for sheet_name, df in sheets.items():
        df = df.fillna("")
        q_col = _find_column(df.columns, QUESTION_HEADER_HINTS)
        if q_col is None:
            continue
        a_col = _find_column(df.columns, ANSWER_HEADER_HINTS)
        for idx, row in df.iterrows():
            text = str(row[q_col]).strip()
            if not text or len(text) < 8:
                continue
            existing = str(row[a_col]).strip() if a_col else ""
            questions.append(
                Question(
                    id=f"{sheet_name}!{q_col}{idx + 2}",  # +2: header row + 0-index
                    sheet=sheet_name,
                    row=int(idx) + 2,
                    text=text,
                    existing_answer=existing,
                )
            )
    return questions


def _header_index(header: list[str], hints) -> int | None:
    for hint in hints:
        for i, h in enumerate(header):
            if hint in h:
                return i
    return None


def _parse_qa_table(table, t_idx: int, q_col: int, a_col: int) -> list[Question]:
    """A table with explicit Question and Answer columns (e.g. "Nº | QUESTION |
    ANSWER"). Every data row is a question, question mark or not; numbered rows
    without a dot ("5") are section headings when the table also has dotted
    numbers ("5.1"). Each question remembers its exact (table, row, answer
    column) so the answer is written into that cell and nowhere else."""
    rows = list(table.rows)[1:]
    numbers = [r.cells[0].text.strip() for r in rows]
    has_sub_numbers = any(NUMBER_RE.match(n) and "." in n.rstrip(".") for n in numbers)

    questions: list[Question] = []
    section = ""
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            continue
        cells = row.cells
        if max(q_col, a_col) >= len(cells) or cells[q_col]._tc is cells[a_col]._tc:
            continue  # row merged across the answer column: a banner, not a question
        text = cells[q_col].text.strip()
        if not text:
            continue
        num = cells[0].text.strip() if q_col != 0 else ""
        if has_sub_numbers and NUMBER_RE.match(num) and "." not in num.rstrip("."):
            section = text
            continue
        questions.append(
            Question(
                id=num or f"table{t_idx}:r{r_idx}",
                sheet=f"table{t_idx}",
                row=r_idx,
                text=text,
                existing_answer=cells[a_col].text.strip(),
                section=section,
                answer_col=a_col,
            )
        )
    return questions


def parse_docx(path: Path) -> list[Question]:
    questions: list[Question] = []
    doc = DocxDocument(str(path))
    counter = 0

    def maybe_add(text: str, sheet: str, row: int):
        nonlocal counter
        text = text.strip()
        if len(text) >= 8 and ("?" in text or text[:2].isdigit()):
            counter += 1
            questions.append(Question(id=f"p{counter}", sheet=sheet, row=row, text=text))

    for i, para in enumerate(doc.paragraphs):
        maybe_add(para.text, sheet="body", row=i)

    for t_idx, table in enumerate(doc.tables):
        header = [c.text.strip().lower() for c in table.rows[0].cells] if table.rows else []
        q_col_idx = _header_index(header, QUESTION_HEADER_HINTS)
        a_col_idx = _header_index(header, ANSWER_HEADER_HINTS)
        if q_col_idx is not None and a_col_idx is not None and q_col_idx != a_col_idx:
            questions.extend(_parse_qa_table(table, t_idx, q_col_idx, a_col_idx))
            continue
        # No Question/Answer header: fall back to the "looks like a question" heuristic.
        # These are reported but can't be written back (no known answer slot).
        for r_idx, row in enumerate(table.rows):
            if q_col_idx is not None:
                if r_idx == 0:
                    continue
                cells = row.cells
                if q_col_idx < len(cells):
                    maybe_add(cells[q_col_idx].text, sheet=f"table{t_idx}", row=r_idx)
            else:
                for cell in row.cells:
                    maybe_add(cell.text, sheet=f"table{t_idx}", row=r_idx)

    return questions


def numbering_gaps(questions: list[Question]) -> list[str]:
    """Spot skipped questions: for dotted ids like 3.1, 3.2, 3.4 report "3.3".
    Cheap sanity check that the parser didn't silently drop a row."""
    seen: dict[str, list[int]] = {}
    for q in questions:
        parts = q.id.rstrip(".").split(".")
        if len(parts) >= 2 and all(p.isdigit() for p in parts):
            seen.setdefault(".".join(parts[:-1]), []).append(int(parts[-1]))
    gaps = []
    for prefix, nums in seen.items():
        for n in range(1, max(nums) + 1):
            if n not in nums:
                gaps.append(f"{prefix}.{n}")
    return gaps


def parse_questionnaire(path: Path) -> list[Question]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Unsupported questionnaire format: {suffix} (use .xlsx or .docx)")


def main() -> int:
    ap = argparse.ArgumentParser(description="Parse a questionnaire and print extracted questions.")
    ap.add_argument("questionnaire", help="Path to .xlsx or .docx questionnaire")
    ap.add_argument("--json", action="store_true", help="Print as JSON instead of a summary")
    args = ap.parse_args()

    path = Path(args.questionnaire)
    questions = parse_questionnaire(path)

    if args.json:
        print(json.dumps([asdict(q) for q in questions], indent=2))
    else:
        placeable = sum(1 for q in questions if q.answer_col is not None or path.suffix.lower() == ".xlsx")
        print(f"Found {len(questions)} questions in {path.name} "
              f"({placeable} with a known answer slot):\n")
        for q in questions[:20]:
            print(f"  [{q.id}] {q.text[:100]}")
        if len(questions) > 20:
            print(f"  ... and {len(questions) - 20} more")
        gaps = numbering_gaps(questions)
        if gaps:
            print(f"\nWARNING: numbering gaps, possibly missed questions: {', '.join(gaps)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
